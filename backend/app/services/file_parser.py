"""Parsing of CSV and XLSX files into one clean, internal CSV format.

Both formats go through the same steps:
    rows from the file -> clean and validate -> write to a CSV file (UTF-8, comma-separated)

The parsers never load the whole table as Python objects. They write row by row,
so the row and column limits can stop a very large file early.
"""

import codecs
import csv
import io
import logging
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time
from itertools import islice
from pathlib import Path
from typing import TextIO

from openpyxl import load_workbook

from app.core.config import BYTES_PER_MB
from app.core.errors import AppError, ErrorCode
from app.utils.table_validator import validate_headers

logger = logging.getLogger(__name__)

NumberedRow = tuple[int, list[str]]  # (line or row number in the source file, cells)


@dataclass(frozen=True)
class TableLimits:
    max_rows: int
    max_columns: int


@dataclass(frozen=True)
class ParseResult:
    columns: list[str]
    row_count: int
    warnings: list[str] = field(default_factory=list)


# ----------------------------------------------------------------------------
# Shared step: clean rows, validate them and write them as CSV
# ----------------------------------------------------------------------------
def _trim_trailing_empty(cells: list[str]) -> list[str]:
    while cells and not cells[-1].strip():
        cells.pop()
    return cells


def _write_table(
    rows: Iterable[NumberedRow],
    destination: TextIO,
    limits: TableLimits,
    *,
    source_format: str,
) -> tuple[list[str], int]:
    """Validate the rows and write them to `destination`. Returns (column names, data row count)."""
    writer = csv.writer(destination, lineterminator="\n")
    columns: list[str] | None = None
    row_count = 0

    for line_number, raw_cells in rows:
        cells = _trim_trailing_empty(list(raw_cells))
        if not cells or not any(cell.strip() for cell in cells):
            continue  # completely blank row: ignore

        if columns is None:  # the first non-blank row is the header
            columns = validate_headers(
                cells,
                max_columns=limits.max_columns,
                check_separator=(source_format == "csv"),
            )
            writer.writerow(columns)
            continue

        if len(cells) > len(columns):
            raise AppError(
                ErrorCode.MALFORMED_ROWS,
                f"Row {line_number} has more values than there are column names. "
                "The first row of the file must contain the column names.",
                details={"row": line_number, "expected_columns": len(columns)},
            )

        row_count += 1
        if row_count > limits.max_rows:
            raise AppError(
                ErrorCode.TOO_MANY_ROWS,
                f"The file has more than {limits.max_rows} data rows, which is the limit.",
                details={"max_rows": limits.max_rows},
            )
        cells.extend([""] * (len(columns) - len(cells)))  # short rows are padded
        writer.writerow(cells)

    if columns is None:
        raise AppError(ErrorCode.EMPTY_FILE, "The file has no content.")
    if row_count == 0:
        raise AppError(
            ErrorCode.NO_DATA_ROWS,
            "The file has column names but no data rows.",
            details={"columns": columns},
        )
    return columns, row_count


# ----------------------------------------------------------------------------
# CSV
# ----------------------------------------------------------------------------
def _decode_csv_text(data: bytes, warnings: list[str]) -> str:
    try:
        return data.decode("utf-8-sig")  # also removes the BOM that Excel adds
    except UnicodeDecodeError:
        pass
    try:
        text = data.decode("cp1252")
    except UnicodeDecodeError as exc:
        raise AppError(
            ErrorCode.INVALID_ENCODING,
            "The file text encoding is not supported. Please save it as 'CSV UTF-8' and upload again.",
        ) from exc
    warnings.append(
        "The file was not UTF-8, so it was read as Windows-1252 (the default of Excel CSV). "
        "Please check that special characters look correct."
    )
    return text


def parse_csv(data: bytes, destination: TextIO, limits: TableLimits) -> ParseResult:
    warnings: list[str] = []

    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        raise AppError(
            ErrorCode.INVALID_ENCODING,
            "UTF-16 files are not supported. Please save the file as 'CSV UTF-8' and upload again.",
        )
    if b"\x00" in data:
        raise AppError(
            ErrorCode.CORRUPTED_FILE,
            "The file does not look like a text CSV file. It may be corrupted or in another format.",
        )

    text = _decode_csv_text(data, warnings)
    reader = csv.reader(io.StringIO(text, newline=""), strict=True)

    def numbered_rows() -> Iterator[NumberedRow]:
        try:
            for cells in reader:
                yield reader.line_num, cells
        except csv.Error as exc:
            raise AppError(
                ErrorCode.MALFORMED_ROWS,
                f"The CSV file could not be read near line {reader.line_num}: {exc}",
                details={"line": reader.line_num},
            ) from exc

    columns, row_count = _write_table(numbered_rows(), destination, limits, source_format="csv")
    return ParseResult(columns, row_count, warnings)


# ----------------------------------------------------------------------------
# XLSX (first worksheet only)
# ----------------------------------------------------------------------------
def _cell_to_text(value: object) -> str:
    """Convert an Excel cell value to text for the CSV file."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, datetime):  # check datetime before date (datetime is a kind of date)
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return str(value)


def _corrupted_xlsx_error() -> AppError:
    return AppError(
        ErrorCode.CORRUPTED_FILE,
        "The .xlsx file could not be read. It may be corrupted, password-protected, or not a real .xlsx file.",
    )


def _check_uncompressed_size(data: bytes, max_bytes: int) -> None:
    """An .xlsx is a zip file. Refuse files that would become huge after unzipping (zip bomb)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(item.file_size for item in archive.infolist())
    except Exception as exc:  # any failure here means the file is not a readable zip
        raise _corrupted_xlsx_error() from exc

    if total > max_bytes:
        raise AppError(
            ErrorCode.FILE_TOO_LARGE,
            f"The workbook is larger than {max_bytes / BYTES_PER_MB:g} MB after unzipping, which is the limit.",
            status_code=413,
            details={"max_uncompressed_bytes": max_bytes},
        )


def _iter_xlsx_rows(sheet: object, max_columns: int) -> Iterator[NumberedRow]:
    try:
        # Read only one column more than allowed. That is enough to detect "too many columns".
        row_iterator = sheet.iter_rows(values_only=True, max_col=max_columns + 1)  # type: ignore[attr-defined]
        for index, row in enumerate(row_iterator, start=1):
            yield index, [_cell_to_text(value) for value in row]
    except Exception as exc:  # broken sheet XML, bad shared strings, etc.
        logger.warning("Could not read worksheet rows (%s)", type(exc).__name__)
        raise _corrupted_xlsx_error() from exc


def parse_xlsx(
    data: bytes,
    destination: TextIO,
    limits: TableLimits,
    max_uncompressed_bytes: int,
) -> ParseResult:
    if not data.startswith(b"PK\x03\x04"):  # every real .xlsx starts with the zip signature
        raise _corrupted_xlsx_error()
    _check_uncompressed_size(data, max_uncompressed_bytes)

    try:
        # read_only streams the sheet. data_only=True gives the last saved value of formulas.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        logger.warning("Could not open workbook (%s)", type(exc).__name__)
        raise _corrupted_xlsx_error() from exc

    try:
        if not workbook.worksheets:
            raise AppError(ErrorCode.NO_DATA_ROWS, "The workbook has no worksheet with data.")

        first_sheet = workbook.worksheets[0]
        warnings: list[str] = []
        if len(workbook.worksheets) > 1:
            warnings.append(
                f"The workbook has {len(workbook.worksheets)} sheets. "
                f"Only the first sheet ('{first_sheet.title[:50]}') was imported."
            )

        rows = _iter_xlsx_rows(first_sheet, limits.max_columns)
        columns, row_count = _write_table(rows, destination, limits, source_format="xlsx")
        return ParseResult(columns, row_count, warnings)
    finally:
        workbook.close()


# ----------------------------------------------------------------------------
# Preview (reads the stored CSV, never the original upload)
# ----------------------------------------------------------------------------
def read_preview(
    csv_path: Path,
    rows: int,
    max_cell_length: int,
) -> tuple[list[list[str | None]], bool]:
    """Return the first `rows` data rows. Empty cells become None. Long cells are shortened."""
    preview: list[list[str | None]] = []
    truncated = False

    with csv_path.open("r", encoding="utf-8", newline="") as file:
        reader = csv.reader(file)
        next(reader, None)  # skip the header row
        for cells in islice(reader, rows):
            row: list[str | None] = []
            for cell in cells:
                if cell == "":
                    row.append(None)
                elif len(cell) > max_cell_length:
                    row.append(cell[: max_cell_length - 1] + "…")
                    truncated = True
                else:
                    row.append(cell)
            preview.append(row)

    return preview, truncated